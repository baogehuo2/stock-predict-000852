from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import time
from dataclasses import asdict, dataclass
from datetime import date, datetime, time as datetime_time, timedelta
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from src.common.network import disable_env_proxies


INDEX_CODE = "000852"
INDEX_NAME = "中证1000"
DATA_SOURCE = "csindex:official_announcement"
API_BASE = "https://www.csindex.com.cn/csindex-home"
ANNOUNCEMENT_LIST_URL = f"{API_BASE}/announcement/queryAnnouncementByVonew"
ANNOUNCEMENT_DETAIL_URL = f"{API_BASE}/announcement/queryAnnouncementById"
DEFAULT_ANNOUNCEMENT_IDS = (3006000, 15690, 14842, 13334, 12466)
DATE_PATTERN = re.compile(
    r"(20(?:\s*\d){2})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日"
)
QUOTED_TITLE_PATTERN = re.compile(r"[“\"]([^”\"]+)[”\"]")


@dataclass(frozen=True)
class Announcement:
    announcement_id: int
    publish_date: date
    title: str
    content_text: str
    source_url: str
    attachment_urls: tuple[str, ...]
    attachment_names: tuple[str, ...]
    implementation_after_close_date: date | None
    direct_effective_date: date | None
    available_time: datetime
    is_cancellation: bool
    cancelled_title: str | None


def _plain_text(html: str) -> str:
    return " ".join(BeautifulSoup(html or "", "html.parser").stripped_strings)


def _extract_attachment_urls(data: dict[str, Any]) -> tuple[tuple[str, ...], tuple[str, ...]]:
    urls: list[str] = []
    names: list[str] = []
    for item in data.get("enclosureList") or []:
        url = str(item.get("fileUrl") or "").strip()
        if url:
            urls.append(url)
            names.append(str(item.get("fileName") or Path(url).name))
    if not urls:
        soup = BeautifulSoup(data.get("content") or "", "html.parser")
        for anchor in soup.find_all("a", href=True):
            url = str(anchor["href"]).strip()
            if url.startswith(("http://", "https://")):
                urls.append(url)
                names.append(anchor.get_text(" ", strip=True) or Path(url).name)
    return tuple(urls), tuple(names)


def _parse_implementation_dates(text: str) -> tuple[date | None, date | None]:
    match = DATE_PATTERN.search(text)
    if not match:
        return None, None
    parsed = date(*(int(re.sub(r"\s+", "", value)) for value in match.groups()))
    vicinity = text[match.start() : match.end() + 12]
    if "收市后生效" in vicinity:
        return parsed, None
    if "起" in vicinity or "生效" in vicinity:
        return None, parsed
    return None, None


def _cancellation_title(title: str, content_text: str) -> str | None:
    if not any(token in title + content_text for token in ("不再实施", "不实施", "取消实施")):
        return None
    match = QUOTED_TITLE_PATTERN.search(content_text)
    return match.group(1).strip() if match else None


def normalize_announcement(data: dict[str, Any]) -> Announcement:
    publish_date = pd.to_datetime(data["publishDate"], errors="raise").date()
    title = str(data.get("title") or "").strip()
    content_text = _plain_text(str(data.get("content") or ""))
    attachment_urls, attachment_names = _extract_attachment_urls(data)
    after_close, direct = _parse_implementation_dates(content_text)
    cancelled_title = _cancellation_title(title, content_text)
    return Announcement(
        announcement_id=int(data["id"]),
        publish_date=publish_date,
        title=title,
        content_text=content_text,
        source_url=f"{ANNOUNCEMENT_DETAIL_URL}?id={int(data['id'])}",
        attachment_urls=attachment_urls,
        attachment_names=attachment_names,
        implementation_after_close_date=after_close,
        direct_effective_date=direct,
        # The API exposes only a date, not a publication timestamp. Using the next
        # calendar day prevents accidental same-day look-ahead in later research.
        available_time=datetime.combine(publish_date + timedelta(days=1), datetime_time.min),
        is_cancellation=cancelled_title is not None,
        cancelled_title=cancelled_title,
    )


def fetch_announcement(announcement_id: int, session: requests.Session | None = None) -> Announcement:
    client = session or requests.Session()
    response = client.get(
        ANNOUNCEMENT_DETAIL_URL,
        params={"id": announcement_id},
        timeout=(15, 60),
    )
    response.raise_for_status()
    payload = response.json()
    if str(payload.get("code")) != "200" or not payload.get("data"):
        raise RuntimeError(f"official announcement API failed for id={announcement_id}: {payload}")
    return normalize_announcement(payload["data"])


def find_cancelled_announcement(
    cancellation: Announcement,
    candidates: Iterable[Announcement],
) -> int | None:
    if not cancellation.is_cancellation or not cancellation.cancelled_title:
        return None
    matches = [
        item
        for item in candidates
        if item.publish_date < cancellation.publish_date
        and cancellation.cancelled_title in item.title
    ]
    if not matches:
        return None
    return max(matches, key=lambda item: (item.publish_date, item.announcement_id)).announcement_id


def resolve_effective_date(
    announcement: Announcement,
    trading_dates: Iterable[date] | None = None,
) -> date | None:
    if announcement.is_cancellation:
        return None
    if announcement.direct_effective_date:
        return announcement.direct_effective_date
    if not announcement.implementation_after_close_date or trading_dates is None:
        return None
    later_dates = sorted(
        item for item in set(trading_dates) if item > announcement.implementation_after_close_date
    )
    return later_dates[0] if later_dates else None


def download_attachment(
    url: str,
    session: requests.Session | None = None,
    retries: int = 3,
    read_timeout: int = 180,
) -> bytes:
    client = session or requests.Session()
    headers = {"User-Agent": "Mozilla/5.0", "Referer": "https://www.csindex.com.cn/"}
    last_error: Exception | None = None
    for attempt in range(1, retries + 1):
        try:
            response = client.get(url, headers=headers, timeout=(20, read_timeout))
            response.raise_for_status()
            content = response.content
            suffix = Path(url.split("?", 1)[0]).suffix.lower()
            if suffix == ".pdf" and not content.startswith(b"%PDF"):
                raise ValueError(f"official PDF attachment has invalid signature: {content[:8]!r}")
            if suffix in {".xlsx", ".xlsm"} and not content.startswith(b"PK"):
                raise ValueError(f"official XLSX attachment has invalid signature: {content[:8]!r}")
            if not content:
                raise ValueError("official attachment is empty")
            return content
        except (requests.RequestException, ValueError) as exc:
            last_error = exc
            if attempt < retries:
                time.sleep(float(attempt))
    raise RuntimeError(f"failed to download official attachment after {retries} attempts: {url}") from last_error


def _normalize_adjustment_rows(rows: list[dict[str, Any]]) -> pd.DataFrame:
    frame = pd.DataFrame(rows)
    columns = [
        "index_code", "index_name", "stock_code", "stock_name", "adjustment_type"
    ]
    if frame.empty:
        return pd.DataFrame(columns=columns)
    frame["index_code"] = frame["index_code"].astype(str).str.extract(r"(\d{6})", expand=False)
    frame["stock_code"] = frame["stock_code"].astype(str).str.extract(r"(\d{6})", expand=False)
    frame["index_name"] = frame["index_name"].astype(str).str.replace(r"\s+", "", regex=True)
    frame["stock_name"] = frame["stock_name"].astype(str).str.replace(r"\s+", "", regex=True)
    frame = frame[
        (frame["index_code"] == INDEX_CODE)
        | frame["index_name"].str.contains(INDEX_NAME, na=False)
    ].copy()
    if frame.empty:
        return pd.DataFrame(columns=columns)
    frame["index_code"] = INDEX_CODE
    frame["index_name"] = INDEX_NAME
    if frame["stock_code"].isna().any() or frame["stock_code"].duplicated().any():
        raise ValueError("invalid or duplicate CSI 1000 stock codes in adjustment attachment")
    return frame[columns].sort_values(["adjustment_type", "stock_code"]).reset_index(drop=True)


def parse_adjustment_xlsx(content: bytes) -> pd.DataFrame:
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    rows: list[dict[str, Any]] = []
    for sheet_name, adjustment_type in (("换出", "remove"), ("换入", "add")):
        if sheet_name not in workbook.sheetnames:
            continue
        values = list(workbook[sheet_name].iter_rows(values_only=True))
        if not values:
            continue
        headers = [str(value or "").replace("\n", "").strip() for value in values[0]]
        required = ("指数代码", "指数简称", "证券代码", "证券简称")
        if not all(name in headers for name in required):
            raise ValueError(f"unexpected {sheet_name} sheet headers: {headers}")
        positions = {name: headers.index(name) for name in required}
        for values_row in values[1:]:
            if not any(value is not None for value in values_row):
                continue
            rows.append(
                {
                    "index_code": values_row[positions["指数代码"]],
                    "index_name": values_row[positions["指数简称"]],
                    "stock_code": values_row[positions["证券代码"]],
                    "stock_name": values_row[positions["证券简称"]],
                    "adjustment_type": adjustment_type,
                }
            )
    return _normalize_adjustment_rows(rows)


def _parse_regular_adjustment_text(text: str) -> pd.DataFrame:
    start_match = re.search(r"中证\s*1000\s*指数样本调整名单[：:]?", text)
    if not start_match:
        raise ValueError("CSI 1000 adjustment section not found in official PDF")
    section = text[start_match.end() :]
    end_match = re.search(r"\n中证\s*(?!1000\b)[A-Za-z0-9]+\s*指数样本调整名单", section)
    if end_match:
        section = section[: end_match.start()]

    rows: list[dict[str, Any]] = []
    row_pattern = re.compile(r"^(\d{6})\s+(.+?)\s+(\d{6})\s+(.+)$")
    for raw_line in section.splitlines():
        match = row_pattern.match(raw_line.strip())
        if not match:
            continue
        remove_code, remove_name, add_code, add_name = match.groups()
        rows.extend(
            [
                {
                    "index_code": INDEX_CODE,
                    "index_name": INDEX_NAME,
                    "stock_code": remove_code,
                    "stock_name": remove_name,
                    "adjustment_type": "remove",
                },
                {
                    "index_code": INDEX_CODE,
                    "index_name": INDEX_NAME,
                    "stock_code": add_code,
                    "stock_name": add_name,
                    "adjustment_type": "add",
                },
            ]
        )
    return _normalize_adjustment_rows(rows)


def parse_adjustment_pdf(content: bytes) -> pd.DataFrame:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(content)) as document:
        text = "\n".join(page.extract_text() or "" for page in document.pages)
    return _parse_regular_adjustment_text(text)


def parse_attachment(url: str, content: bytes) -> pd.DataFrame:
    suffix = Path(url.split("?", 1)[0]).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return parse_adjustment_xlsx(content)
    if suffix == ".pdf":
        return parse_adjustment_pdf(content)
    raise ValueError(f"unsupported official attachment type: {suffix}")


def validate_announcements(
    announcement_ids: Iterable[int] = DEFAULT_ANNOUNCEMENT_IDS,
    download_files: bool = True,
) -> dict[str, Any]:
    disable_env_proxies()
    session = requests.Session()
    announcements = [fetch_announcement(item, session=session) for item in announcement_ids]
    cancellations = {
        item.announcement_id: find_cancelled_announcement(item, announcements)
        for item in announcements
        if item.is_cancellation
    }
    cancelled_targets = {item for item in cancellations.values() if item is not None}
    parsed_files: list[dict[str, Any]] = []
    skipped_files: list[dict[str, Any]] = []
    if download_files:
        for announcement in announcements:
            if announcement.is_cancellation:
                continue
            if announcement.announcement_id in cancelled_targets:
                skipped_files.append(
                    {
                        "announcement_id": announcement.announcement_id,
                        "reason": "cancelled_by_later_official_announcement",
                    }
                )
                continue
            for name, url in zip(announcement.attachment_names, announcement.attachment_urls):
                content = download_attachment(url, session=session)
                frame = parse_attachment(url, content)
                add_rows = int((frame["adjustment_type"] == "add").sum())
                remove_rows = int((frame["adjustment_type"] == "remove").sum())
                if not len(frame) or add_rows != remove_rows:
                    raise ValueError(
                        f"invalid adjustment counts for announcement "
                        f"{announcement.announcement_id}: add={add_rows}, remove={remove_rows}"
                    )
                if Path(url.split("?", 1)[0]).suffix.lower() == ".pdf" and add_rows != 100:
                    raise ValueError(
                        f"regular CSI 1000 adjustment must contain 100 adds and removes, "
                        f"got add={add_rows}, remove={remove_rows}"
                    )
                parsed_files.append(
                    {
                        "announcement_id": announcement.announcement_id,
                        "attachment_name": name,
                        "source_url": url,
                        "sha256": hashlib.sha256(content).hexdigest(),
                        "rows": len(frame),
                        "add_rows": add_rows,
                        "remove_rows": remove_rows,
                        "sample": frame.head(5).to_dict(orient="records"),
                    }
                )
    return {
        "index_code": INDEX_CODE,
        "data_source": DATA_SOURCE,
        "history_write_enabled": False,
        "announcements": [asdict(item) for item in announcements],
        "cancellations": cancellations,
        "cancelled_announcement_ids": sorted(cancelled_targets),
        "skipped_files": skipped_files,
        "parsed_files": parsed_files,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="验证中证1000官方历史调样公告，不写入历史表")
    parser.add_argument("--announcement-id", action="append", type=int, dest="announcement_ids")
    parser.add_argument("--metadata-only", action="store_true")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    report = validate_announcements(
        announcement_ids=args.announcement_ids or DEFAULT_ANNOUNCEMENT_IDS,
        download_files=not args.metadata_only,
    )
    payload = json.dumps(report, ensure_ascii=False, indent=2, default=str)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(payload + "\n", encoding="utf-8")
    print(payload)


if __name__ == "__main__":
    main()
